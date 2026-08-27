"""
Compute Cohen's kappa and Gwet's AC1 for ALL four manual-evaluation studies.

Files:
  01_usefulness_pipelines.xlsx   -- 62 queries x 11 retrieval methods
  02_usefulness_thresholds.xlsx  -- 62 queries x HB1 x 9 thresholds
  03_final_answer_usefulness.xlsx-- 68 queries x 6 generators (developer survey)
  04_hallucination.xlsx       -- 68-row hallucination alignment sample

All four use the same column convention: `annotator2_label` and
`annotator3_label`, binary labels.
"""
import openpyxl
from sklearn.metrics import cohen_kappa_score


def norm(v):
    return str(v).strip().lower() if v not in (None, "") else None


def gwet_ac1_binary(a, b, positive):
    n = len(a)
    p_a = sum(1 for x, y in zip(a, b) if x == y) / n
    pi1 = sum(1 for x in a if x == positive) / n
    pi2 = sum(1 for x in b if x == positive) / n
    pi_bar = (pi1 + pi2) / 2
    p_e = 2 * pi_bar * (1 - pi_bar)
    ac1 = (p_a - p_e) / (1 - p_e) if p_e != 1 else float("nan")
    return ac1, p_a, pi1, pi2


def label_interp(k):
    if k != k:
        return "n/a"
    if k < 0.20: return "poor"
    if k < 0.40: return "fair"
    if k < 0.60: return "moderate"
    if k < 0.80: return "substantial"
    return "almost-perfect"


def process(path, file_label, positive=None):
    print(f"\n========== {file_label}: {path} ==========")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    print(f"{'sheet':<28}{'n':>6}{'raw%':>8}"
          f"{'kappa':>9}{'k-tag':<14}{'AC1':>9}{'AC1-tag':<14}"
          f"{'pi2':>8}{'pi3':>8}")
    all_a, all_b = [], []
    auto_pos = positive
    for sh in wb.sheetnames:
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        hdr = rows[0]
        if "annotator2_label" not in hdr or "annotator3_label" not in hdr:
            continue
        i2 = hdr.index("annotator2_label")
        i3 = hdr.index("annotator3_label")
        a, b = [], []
        for r in rows[1:]:
            if r[0] is None:
                continue
            v2, v3 = norm(r[i2]), norm(r[i3])
            if v2 is None or v3 is None:
                continue
            a.append(v2); b.append(v3)
        if not a:
            continue
        # Detect positive label automatically if not specified
        if auto_pos is None:
            counts = {}
            for x in a + b:
                counts[x] = counts.get(x, 0) + 1
            auto_pos = max(counts, key=counts.get)
        k = cohen_kappa_score(a, b)
        ac1, raw, pi1, pi2 = gwet_ac1_binary(a, b, auto_pos)
        print(f"{sh:<28}{len(a):>6}{raw*100:>7.1f}%"
              f"{k:>9.3f}  {label_interp(k):<12}"
              f"{ac1:>9.3f}  {label_interp(ac1):<12}"
              f"{pi1:>8.3f}{pi2:>8.3f}")
        all_a += a; all_b += b
    if all_a:
        print("-" * 92)
        k = cohen_kappa_score(all_a, all_b)
        ac1, raw, pi1, pi2 = gwet_ac1_binary(all_a, all_b, auto_pos)
        print(f"{'ALL':<28}{len(all_a):>6}{raw*100:>7.1f}%"
              f"{k:>9.3f}  {label_interp(k):<12}"
              f"{ac1:>9.3f}  {label_interp(ac1):<12}"
              f"{pi1:>8.3f}{pi2:>8.3f}")


if __name__ == "__main__":
    process("01_usefulness_pipelines.xlsx",
            "File 01 -- Retrieval Usefulness across 11 methods (62q x 11)")
    process("02_usefulness_thresholds.xlsx",
            "File 02 -- HB1 across 9 thresholds (62q x 9 thr)")
    process("03_final_answer_usefulness.xlsx",
            "File 03 -- Developer survey of final-answer helpfulness (68q x 6 models = 408)")
    process("04_hallucination.xlsx",
            "File 04 -- Hallucination 68-row LLM-vs-human alignment")
