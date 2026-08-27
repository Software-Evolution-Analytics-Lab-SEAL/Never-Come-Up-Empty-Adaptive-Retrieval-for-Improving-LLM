"""
Compute Gwet's AC1 alongside Cohen's kappa for the threshold-sweep manual
evaluation (file 02_usefulness_thresholds.xlsx).

Background:
- Cohen's kappa is paradoxically deflated when one class dominates (here:
  most retrieved units are labelled "useful"), even though raw agreement is
  high. Gwet's AC1 (Gwet 2008) was designed to be robust to this prevalence
  skew.

Gwet's AC1 (binary):
    P_e_gwet = 2 * pi_bar * (1 - pi_bar)
where pi_bar is the marginal "useful" rate averaged across both annotators.
    AC1 = (P_a - P_e_gwet) / (1 - P_e_gwet)

Run from the project root:
    cd New_Experiments_Code/Manual_Evaluation_Results
    python3 compute_gwet_ac1_threshold.py
"""
import openpyxl
import pandas as pd
from sklearn.metrics import cohen_kappa_score


def gwet_ac1_binary(a, b, positive="true"):
    """Gwet's AC1 for two raters on a binary task.

    a, b -- iterables of equal length containing string labels.
    positive -- the label considered the "positive" class.
    """
    a = [str(x).strip().lower() for x in a]
    b = [str(x).strip().lower() for x in b]
    pos = positive.strip().lower()
    n = len(a)
    p_a = sum(1 for x, y in zip(a, b) if x == y) / n
    pi1 = sum(1 for x in a if x == pos) / n
    pi2 = sum(1 for x in b if x == pos) / n
    pi_bar = (pi1 + pi2) / 2
    p_e = 2 * pi_bar * (1 - pi_bar)
    return (p_a - p_e) / (1 - p_e) if p_e != 1 else float("nan"), p_a, pi1, pi2


def main():
    path = "02_usefulness_thresholds.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    all_a2, all_a3 = [], []
    print(f"{'sheet':<10}{'n':>6}{'raw%':>8}"
          f"{'kappa':>9}{'AC1':>9}{'pi1':>8}{'pi2':>8}")
    for sh in wb.sheetnames:
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        hdr = rows[0]
        if "annotator2_label" not in hdr or "annotator3_label" not in hdr:
            continue
        i2 = hdr.index("annotator2_label")
        i3 = hdr.index("annotator3_label")
        a2, a3 = [], []
        for r in rows[1:]:
            if r[0] is None:
                continue
            v2, v3 = r[i2], r[i3]
            if v2 in (None, "") or v3 in (None, ""):
                continue
            a2.append(str(v2).strip().lower())
            a3.append(str(v3).strip().lower())
        if not a2:
            continue
        k = cohen_kappa_score(a2, a3)
        ac1, raw, pi1, pi2 = gwet_ac1_binary(a2, a3, positive="true")
        print(f"{sh:<10}{len(a2):>6}{raw*100:>7.1f}%"
              f"{k:>9.3f}{ac1:>9.3f}{pi1:>8.3f}{pi2:>8.3f}")
        all_a2 += a2
        all_a3 += a3

    print("-" * 56)
    n = len(all_a2)
    k = cohen_kappa_score(all_a2, all_a3)
    ac1, raw, pi1, pi2 = gwet_ac1_binary(all_a2, all_a3, positive="true")
    print(f"{'ALL':<10}{n:>6}{raw*100:>7.1f}%"
          f"{k:>9.3f}{ac1:>9.3f}{pi1:>8.3f}{pi2:>8.3f}")


if __name__ == "__main__":
    main()
